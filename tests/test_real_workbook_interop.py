"""Round-trip a workbook written by openpyxl, not by our own test helpers.

The optimization that skips re-serializing untouched worksheets is only safe if
the resulting package still opens in tooling that follows the OOXML spec. These
tests write the input with ``openpyxl``, process it, and then read both the
encrypted and the restored files back with ``openpyxl`` as well, so a
spec-violating part would fail here rather than in a client's Excel.

Skipped when ``openpyxl`` is not installed; it is a development convenience, not
a runtime dependency of the application.
"""

import secrets
import tempfile
import unittest
import zipfile
from pathlib import Path

from processor import CellCipher, transform_workbook

try:
    import openpyxl
except ImportError:
    openpyxl = None

SHEET_NAMES = ("Clients", "Notes", "Archive")
HEADERS = ("Full Name", "Email", "City")


def _build_workbook(path: Path, rows: int = 5) -> dict[str, list[list[str]]]:
    """Write a multi-sheet workbook and return its plain values by sheet name."""
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    expected: dict[str, list[list[str]]] = {}

    for index, name in enumerate(SHEET_NAMES):
        sheet = workbook.create_sheet(title=name)
        values: list[list[str]] = [list(HEADERS)]
        for row in range(2, rows + 2):
            values.append([f"{header} {name} {row}" for header in HEADERS])
        for value_row in values:
            sheet.append(value_row)
        expected[name] = values

    workbook.save(path)
    return expected


def _read_workbook(path: Path) -> dict[str, list[list[str]]]:
    """Read every cell back out, so a malformed part raises here."""
    workbook = openpyxl.load_workbook(path)
    try:
        return {
            sheet.title: [
                ["" if cell is None else str(cell) for cell in row]
                for row in sheet.iter_rows(values_only=True)
            ]
            for sheet in workbook.worksheets
        }
    finally:
        workbook.close()


def _part_bytes(path: Path, prefix: str) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as archive:
        return {
            name: archive.read(name)
            for name in archive.namelist()
            if name.startswith(prefix)
        }


class RealWorkbookInteropTest(unittest.TestCase):
    def setUp(self) -> None:
        if openpyxl is None:
            self.skipTest("openpyxl is not installed")

    def test_encrypt_and_decrypt_round_trip_a_real_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "clients.xlsx"
            encrypted = root / "clients.encrypted.xlsx"
            restored = root / "clients.restored.xlsx"
            expected = _build_workbook(source)

            cipher = CellCipher(secrets.token_bytes(64))
            stats = transform_workbook(
                source,
                encrypted,
                cipher,
                "encrypt",
                selected_columns={"0|A"},
            )
            self.assertGreater(stats.total, 0)
            self.assertEqual(
                [header for _sheet, header in stats.encrypted_columns], ["Full Name"]
            )

            # The encrypted package must still be readable by a spec-compliant
            # reader, and only column A of the first sheet may have changed.
            read_back = _read_workbook(encrypted)

            encrypted_value = read_back[SHEET_NAMES[0]][1][0]
            self.assertTrue(cipher.is_encrypted(encrypted_value))
            self.assertNotEqual(encrypted_value, expected[SHEET_NAMES[0]][1][0])
            # Column B was not selected, so it must still hold plain text.
            self.assertEqual(
                read_back[SHEET_NAMES[0]][1][1], expected[SHEET_NAMES[0]][1][1]
            )
            self.assertEqual(read_back[SHEET_NAMES[0]][0], list(HEADERS))

            for name in SHEET_NAMES[1:]:
                self.assertEqual(read_back[name], expected[name])

            # Sheets holding nothing selected stay byte-identical.
            before = _part_bytes(source, "xl/worksheets/")
            after = _part_bytes(encrypted, "xl/worksheets/")
            self.assertEqual(sorted(before), sorted(after))
            self.assertNotEqual(
                before["xl/worksheets/sheet1.xml"], after["xl/worksheets/sheet1.xml"]
            )
            self.assertEqual(
                before["xl/worksheets/sheet2.xml"], after["xl/worksheets/sheet2.xml"]
            )
            self.assertEqual(
                before["xl/worksheets/sheet3.xml"], after["xl/worksheets/sheet3.xml"]
            )
            self.assertEqual(
                _part_bytes(source, "[Content_Types].xml"),
                _part_bytes(encrypted, "[Content_Types].xml"),
            )

            transform_workbook(encrypted, restored, cipher, "decrypt")
            self.assertEqual(_read_workbook(restored), expected)

    def test_replace_only_round_trips_a_real_workbook(self) -> None:
        from processor import ReplacementRule

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "clients.xlsx"
            output = root / "clients.replaced.xlsx"
            expected = _build_workbook(source)

            transform_workbook(
                source,
                output,
                None,
                "replace",
                replacements=[
                    ReplacementRule(find="Email Clients 2", replace="CLIENT_001")
                ],
            )

            read_back = _read_workbook(output)
            # Only the first sheet contains the matched text, and the match lives
            # in the shared-string table rather than in the sheet part.
            self.assertEqual(read_back[SHEET_NAMES[0]][1][1], "CLIENT_001")
            self.assertNotEqual(expected[SHEET_NAMES[0]][1][1], "CLIENT_001")
            self.assertEqual(
                read_back[SHEET_NAMES[1]][1][1], expected[SHEET_NAMES[1]][1][1]
            )
            self.assertEqual(read_back[SHEET_NAMES[0]][0], list(HEADERS))
            self.assertEqual(read_back[SHEET_NAMES[2]], expected[SHEET_NAMES[2]])


if __name__ == "__main__":
    unittest.main()
